
from flask import Flask, jsonify, request
import psycopg2
import json
from docx import Document
from fpdf import FPDF
import base64
from datetime import datetime

# Connect to PostgreSQL
conn = psycopg2.connect(
    dbname="poc", user="postgres", password="xFl201X20FabSN2024Nov", host="localhost", port="5432"
)


def fetch_speaker_data():
    with conn.cursor() as cursor:
        cursor.execute("SELECT * FROM Speakers")
        return cursor.fetchall()
# Step 2: Fetch data from the SQL table
query = "SELECT * from event"
with conn.cursor() as cursor:   
    cursor.execute(query)
    rows = cursor.fetchall()

# Step 3: Prepare the PDF
class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'SQL Table Export', align='C', ln=1)

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Page {self.page_no()}', align='C')

pdf = PDF()
pdf.add_page()
pdf.set_font('Arial', '', 10)

# Step 4: Add data to the PDF
for row in rows:
    event_id = row[0]
    even_name = row[1]
    logo = base64.b64encode(row[2]).decode('utf-8') if row[2] else "N/A"  # Convert bytea to base64 string
    tagline = row[3]
    event_date = row[4].strftime('%Y-%m-%d') if isinstance(row[4], datetime) else "N/A"
    abstract = row[5]

    pdf.cell(0, 10, f"Integer: {event_id}, String: {even_name}, Date: {event_date}", ln=1)
    pdf.multi_cell(0, 10, f"Text: {abstract}")
    pdf.cell(0, 10, f"Bytea (Base64): {logo}", ln=1)
    pdf.cell(0, 10, '-' * 80, ln=1)

# Step 5: Save the PDF
pdf.output("sql_table_export.pdf")

# Close the database connection
cursor.close()
conn.close()
